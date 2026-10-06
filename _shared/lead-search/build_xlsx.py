#!/usr/bin/env python3
"""Write reviewed leads to the lead-sheet format.

Usage: python3 build_xlsx.py review.json leads.xlsx
Uses each lead's "email"/"email_source" (pick from "emails_found"; leave empty if none is the creator's own)
and "notes". Columns match the original lead sheet plus dates, email source and notes.
"""
import json, sys
import openpyxl
from openpyxl.styles import Font, PatternFill

leads = sorted(json.load(open(sys.argv[1])), key=lambda x: x['followers'])
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Leads'
ws.append(['Name', 'User Handle', 'Profile link', 'Follower', 'Email', 'Youtube link', 'Location',
           'Last IG post', 'Last YT upload', 'Email source', 'Notes'])
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78')
link = Font(color='0563C1', underline='single')
for l in leads:
    ws.append([l['name'], '@' + l['handle'], f"https://www.instagram.com/{l['handle']}/", l['followers'],
               l.get('email', ''), l.get('youtube'), l.get('location'), l.get('last_ig_post'),
               l.get('last_yt_upload'), l.get('email_source', ''), l.get('notes', '')])
    row = ws[ws.max_row]
    row[3].number_format = '#,##0'
    for cell, target in ((row[2], row[2].value), (row[5], row[5].value), (row[4], row[4].value and 'mailto:' + row[4].value)):
        if target:
            cell.hyperlink = target; cell.font = link
    if not row[4].value:
        row[4].fill = PatternFill('solid', fgColor='FFF2CC')
for col, w in zip('ABCDEFGHIJK', [28, 18, 42, 10, 30, 44, 16, 13, 14, 40, 55]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'
wb.save(sys.argv[2])
print(f'{len(leads)} leads -> {sys.argv[2]}')
